"""Live verification of the savings sweep, the statements and the balance-sheet items.

Waits for the new build, then checks that the backfill produced the expected
numbers on the real database. Creates and removes only its own temporary
asset/liability records.

Run:  FF_ADMIN_PASSWORD=... python tests_live_savings.py
"""
import io
import json
import os
import time
import urllib.error
import urllib.request

B = os.getenv("FF_BASE", "https://familyfinancials.onrender.com")
PASSWORD = os.environ["FF_ADMIN_PASSWORD"]
TARGET_VERSION = os.getenv("FF_VERSION", "1.4.0")

if os.getenv("FF_CONFIRM_PRODUCTION") != "1":
    print("\n" + "!" * 74)
    print("!  LIVE / PRODUCTION database warning")
    print(f"!  This script talks to the REAL database at: {B}")
    print("!  It only touches its own temporary records, but run it with care.")
    print("!  Set FF_CONFIRM_PRODUCTION=1 to suppress this banner.")
    print("!" * 74 + "\n")

passed = failed = 0


def check(name, ok, detail=""):
    global passed, failed
    if ok:
        passed += 1
    else:
        failed += 1
    print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f"  [{detail}]" if detail else ""), flush=True)


def call(method, path, body=None, token=None, raw=False, tries=4, timeout=120):
    last = None
    for _ in range(tries):
        data = None
        headers = {}
        if body is not None:
            data = json.dumps(body).encode()
            headers["Content-Type"] = "application/json"
        if token:
            headers["Authorization"] = "Bearer " + token
        req = urllib.request.Request(B + path, data=data, method=method, headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                payload = r.read()
                if raw:
                    return r.status, payload
                text = payload.decode(errors="replace")
                try:
                    return r.status, json.loads(text)
                except Exception:
                    return r.status, text
        except urllib.error.HTTPError as e:
            payload = e.read()
            if raw:
                return e.code, payload
            text = payload.decode(errors="replace")
            try:
                return e.code, json.loads(text)
            except Exception:
                return e.code, text
        except Exception as exc:
            last = exc
            time.sleep(3)
    return None, f"network error: {last!r}"


def version_at_least(actual, want):
    """True when the running build is `want` or newer."""
    try:
        return tuple(int(x) for x in str(actual).split(".")) >= tuple(int(x) for x in str(want).split("."))
    except (TypeError, ValueError):
        return False


def close(a, b, tol=2.0):
    return abs(float(a) - float(b)) <= tol


print(f"waiting for version {TARGET_VERSION} (or newer) to go live...")
for attempt in range(20):
    s, hz = call("GET", "/healthz")
    ver = hz.get("version") if isinstance(hz, dict) else None
    if version_at_least(ver, TARGET_VERSION):
        print(f"  build {ver} is live (checked {attempt + 1}x)\n")
        break
    print(f"  attempt {attempt + 1}: still {ver}")
    time.sleep(20)
else:
    print("!! new build never appeared - stopping")
    raise SystemExit(1)

check("database still healthy after the schema upgrade",
      hz.get("database_ok") is True and hz.get("backend") == "postgresql", str(hz.get("backend")))
check("storage still persistent", hz.get("storage_persistent") is True)

s, tok = call("POST", "/api/auth/login", {"username": "admin", "password": PASSWORD})
check("admin login", s == 200, f"HTTP {s}")
if s != 200:
    raise SystemExit(1)
TOK = tok["access_token"]

print("\n== dashboard ==")
s, dash = call("GET", "/api/dashboard", token=TOK)
check("dashboard loads", s == 200, f"HTTP {s}")
check("'balance' KPI is gone", "balance" not in dash)
count = dash["transaction_count"]
print(f"     transactions={count}  income={dash['total_income']:,.0f}  "
      f"expenses={dash['total_expenses']:,.0f}  savings={dash['total_savings']:,.0f}")
check("the 44 sweep rows and the loan row are present", count == 887 + 44 + 1,
      f"{count} (expected 932 = 887 original + 44 sweep + 1 loan)")
surplus = dash["total_income"] - dash["total_expenses"]
check("the dashboard reports the savings BALANCE, not lifetime savings",
      close(dash["total_savings"], 526_628_193, 5) and close(dash["cash_balance"], 69_981_538, 5),
      f"savings {dash['total_savings']:,.0f} cash {dash.get('cash_balance', 0):,.0f} surplus {surplus:,.0f}")

print("\n== savings summary ==")
s, summ = call("GET", "/api/savings/summary?months=60", token=TOK)
check("savings summary loads", s == 200 and "months" in summ, f"HTTP {s}")
closed = [m for m in summ["months"] if m["closed"]]
check("44 closed months reported", len(closed) == 44, str(len(closed)))
check("every closed month: the movement into savings is its surplus (plus loan-funded transfers)",
      all(close(m["savings_net"], m["surplus"] + m["savings_in_loan"]) for m in closed),
      str([m["month"] for m in closed
           if not close(m["savings_net"], m["surplus"] + m["savings_in_loan"])][:5]))
check("every closed month carries a sweep row",
      all(m["swept"] for m in closed),
      f"{sum(1 for m in closed if m['swept'])}/{len(closed)}")
check("the savings balance is opening 260M + surplus + the loan-funded 230M",
      close(summ["totals"]["closing_balance"], 526_628_193, 5),
      f"{summ['totals']['closing_balance']:,.0f}")
check("2026-08 keeps the real 230M transfer and sweeps only the month's surplus",
      close([m for m in summ["months"] if m["month"] == "2026-08"][0]["savings_in_loan"], 230_018_462, 5)
      and close([m for m in summ["months"] if m["month"] == "2026-08"][0]["sweep_in"], 539_233, 5),
      "loan-funded 230,018,462 / sweep in 539,233")
check("members are listed for the per-member table",
      isinstance(summ["members"], list) and len(summ["members"]) >= 2, str(summ["members"]))

print("\n== sweep rows are visible as ordinary transactions ==")
s, tx = call("GET", "/api/transactions?limit=5000", token=TOK)
offsets = [t for t in tx if t.get("auto_offset_month")]
check("44 rows are marked as automatic offsets", len(offsets) == 44, str(len(offsets)))
check("offset rows are Savings transfers or Withdrawals",
      all(t["type"] in ("Savings", "Withdrawal") for t in offsets),
      str(sorted({t["type"] for t in offsets})))
check("offset descriptions name their month",
      all(t["auto_offset_month"] in (t.get("description") or "") for t in offsets),
      offsets[0]["description"] if offsets else "")

print("\n== income statement ==")
s, rep = call("GET", "/api/statements/income", token=TOK)
check("income statement loads", s == 200 and "income_lines" in rep, f"HTTP {s}")
check("totals internally consistent", close(rep["income_total"] - rep["expense_total"], rep["surplus"]),
      f"surplus {rep['surplus']:,.0f}")
check("borrowing is never income",
      close(rep["income_total"], 2_054_774_297, 5), f"{rep['income_total']:,.0f}")
check("transfers and loans stay below the line, out of the surplus",
      close(rep["financing"]["savings_in"], 361_374_438, 5)
      and close(rep["financing"]["withdrawals"], 324_764_707, 5)
      and close(rep["financing"]["loan_borrowed"], 300_000_000, 5)
      and close(rep["financing"]["savings_in_funded_by_loan"], 230_018_462, 5),
      f"in {rep['financing']['savings_in']:,.0f} out {rep['financing']['withdrawals']:,.0f}"
      f" borrowed {rep['financing']['loan_borrowed']:,.0f}")
check("income-funded savings net of withdrawals equals the surplus",
      close(rep["financing"]["savings_net"], rep["surplus"], 5),
      f"net {rep['financing']['savings_net']:,.0f} vs surplus {rep['surplus']:,.0f}")
check("expense groups are subtotalled", len(rep["expense_groups"]) >= 2, str(len(rep["expense_groups"])))
check("monthly rows returned", len(rep["months"]) == 44, str(len(rep["months"])))
s, one = call("GET", "/api/statements/income?start=2026-08-01&end=2026-08-31", token=TOK)
check("a single month can be reported", s == 200 and len(one["months"]) == 1,
      f"surplus {one['surplus']:,.0f}")

print("\n== balance sheet ==")
s, bs = call("GET", "/api/statements/balance-sheet", token=TOK)
check("balance sheet loads", s == 200 and "net_worth" in bs, f"HTTP {s}")
check("the balance sheet reconciles: opening balances + lifetime surplus = net worth",
      close(bs["reconciliation"]["difference"], 0, 0.5),
      f"expected {bs['reconciliation']['expected_net_worth']:,.0f} actual {bs['net_worth']:,.0f}")
check("money in the accounts is opening + surplus + what is still borrowed",
      close(bs["money_total"], 260_000_000 + 36_609_731 + 300_000_000, 5),
      f"{bs['money_total']:,.0f}")
base_assets = bs["total_assets"]
base_net = bs["net_worth"]
print(f"     assets={base_assets:,.0f}  liabilities={bs['liabilities_total']:,.0f}  net worth={base_net:,.0f}")

s, a = call("POST", "/api/assets", {"name": "TEMP live check", "kind": "property", "value": 1_000_000_000}, token=TOK)
check("an asset can be added live", s == 201, f"HTTP {s}")
s, l = call("POST", "/api/liabilities", {"name": "TEMP live check loan", "kind": "mortgage",
                                         "outstanding": 400_000_000}, token=TOK)
check("a liability can be added live", s == 201, f"HTTP {s}")
s, bs2 = call("GET", "/api/statements/balance-sheet", token=TOK)
check("the balance sheet picked them up",
      close(bs2["total_assets"], base_assets + 1_000_000_000)
      and close(bs2["net_worth"], base_net + 600_000_000),
      f"assets {bs2['total_assets']:,.0f} net {bs2['net_worth']:,.0f}")
s, _ = call("GET", f"/api/assets?include_inactive=true", token=TOK)
check("assets are listed", s == 200 and any(x["name"] == "TEMP live check" for x in _), str(len(_)))
call("DELETE", f"/api/assets/{a['id']}", token=TOK)
call("DELETE", f"/api/liabilities/{l['id']}", token=TOK)
s, bs3 = call("GET", "/api/statements/balance-sheet", token=TOK)
check("cleanup restored the original net worth", close(bs3["net_worth"], base_net) and not bs3["asset_items"],
      f"{bs3['net_worth']:,.0f}")

print("\n== reports, exports and the reminder survive ==")
s, rep_sum = call("POST", "/api/reports/summary", {"group_by": "month"}, token=TOK)
check("the old reports endpoint still works", s == 200 and rep_sum["rows"], f"{len(rep_sum.get('rows', []))} rows")
s, csvb = call("GET", "/api/export/csv", token=TOK, raw=True)
head = csvb.decode("utf-8-sig").splitlines()[0] if s == 200 else ""
check("CSV export keeps working with the new columns",
      s == 200 and "Member" in head and "AutoOffset" in head, head)
check("CSV holds every row including sweeps", s == 200 and len(csvb.decode('utf-8-sig').splitlines()) - 1 == count,
      f"{len(csvb.decode('utf-8-sig').splitlines()) - 1} vs {count}")
s, xlsx = call("GET", "/api/export/excel", token=TOK, raw=True)
from openpyxl import load_workbook
wb = load_workbook(io.BytesIO(xlsx))
ws = wb["Transactions"]
check("Excel export has the 10 columns and all rows",
      [c.value for c in ws[1]] == ["Date", "Type", "Category", "Amount", "Description", "Member",
                                  "AutoOffset", "Direction", "Lender", "FundedBy"]
      and ws.max_row - 1 == count,
      f"cols={len([c.value for c in ws[1]])}, rows={ws.max_row - 1}")
s, st = call("GET", "/api/admin/backup-status", token=TOK)
check("weekly reminder bookkeeping intact", s == 200 and st["interval_days"] == 7, f"due={st['due']}")
check("status reports the new app version", version_at_least(st.get("app_version"), TARGET_VERSION), str(st.get("app_version")))

print(f"\n===== LIVE SAVINGS CHECK: {passed} passed, {failed} failed =====")
raise SystemExit(1 if failed else 0)
