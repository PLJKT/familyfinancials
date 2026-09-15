"""End-to-end test of the savings sweep, the statements and the balance-sheet items.

Run against a scratch server:
  DATABASE_URL="sqlite:///./test_import.db" python -m uvicorn app.main:app --port 8002
  python tests_statements.py
"""
import io
import json
import time
import urllib.error
import urllib.request

B = "http://127.0.0.1:8002"
passed = failed = 0


def check(name, ok, detail=""):
    global passed, failed
    if ok:
        passed += 1
    else:
        failed += 1
    line = f"  {'PASS' if ok else 'FAIL'}  {name}"
    if detail:
        line += f"  [{detail}]"
    print(line, flush=True)


def call(method, path, body=None, token=None, raw=False, tries=3):
    last = None
    for _ in range(tries):
        data = None
        headers = {}
        if body is not None:
            data = json.dumps(body).encode()
            headers["Content-Type"] = "application/json"
        if token:
            headers["Authorization"] = "Bearer " + token
        req = urllib.request.Request(B + path, data=data, method=method, headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                payload = r.read()
                if raw:
                    return r.status, payload
                text = payload.decode(errors="replace")
                try:
                    return r.status, json.loads(text)
                except Exception:
                    return r.status, text
        except urllib.error.HTTPError as e:
            payload = e.read()
            if raw:
                return e.code, payload
            text = payload.decode(errors="replace")
            try:
                return e.code, json.loads(text)
            except Exception:
                return e.code, text
        except Exception as exc:
            last = exc
            time.sleep(2)
    return None, f"network error: {last!r}"


def close(a, b, tol=2.0):
    return abs(float(a) - float(b)) <= tol


# ---------------------------------------------------------------------- login
s, tok = call("POST", "/api/auth/login", {"username": "admin", "password": "admin123"})
check("admin login", s == 200, f"HTTP {s}")
if s != 200:
    print("cannot continue:", str(tok)[:200])
    raise SystemExit(1)
TOK = tok["access_token"]

print("\n== 1. dashboard no longer shows a balance ==")
s, dash = call("GET", "/api/dashboard", token=TOK)
check("dashboard loads", s == 200, f"HTTP {s}")
check("'balance' key is gone", "balance" not in dash, str(sorted(dash.keys()))[:80])
swept_savings = float(dash["total_savings"])
surplus = float(dash["total_income"]) - float(dash["total_expenses"])
check("total savings now equals income - expenses (the sweep ran)",
      close(swept_savings, surplus), f"{swept_savings:,.0f} vs {surplus:,.0f}")
check("dashboard still reports transactions", dash["transaction_count"] >= 887, str(dash["transaction_count"]))

print("\n== 2. savings summary ==")
s, summary = call("GET", "/api/savings/summary?months=60", token=TOK)
check("savings summary loads", s == 200 and "months" in summary, f"HTTP {s}")
months = summary["months"]
closed = [m for m in months if m["closed"]]
check("44 closed months are reported", len(closed) == 44, f"{len(closed)} closed of {len(months)}")
bad = [m["month"] for m in closed if not close(m["savings_total"], m["surplus"])]
check("every closed month's savings equals its surplus", not bad, f"offending: {bad[:5]}")
offsets = [m for m in closed if m["offset_applied"]]
check("closed months carry an automatic offset", len(offsets) == len(closed), f"{len(offsets)}/{len(closed)}")
running = [m for m in months if not m["closed"]]
check("the running month has no automatic offset",
      all(not m["offset_applied"] for m in running), f"{len(running)} running month(s)")
check("members are listed", isinstance(summary["members"], list) and summary["members"], str(summary["members"]))
check("offset for 2026-08 reflects the 230M lump-sum deposit",
      close([m for m in months if m["month"] == "2026-08"][0]["savings_offset"], -229_479_229, 5),
      str([m for m in months if m["month"] == "2026-08"][0]["savings_offset"]))

print("\n== 3. the sweep is idempotent ==")
s, run1 = call("POST", "/api/admin/offsets/run", token=TOK)
check("manual re-run creates nothing", s == 200 and run1["created"] == 0 and run1["updated"] == 0,
      f"created={run1.get('created')} updated={run1.get('updated')} removed={run1.get('removed')}")

print("\n== 4. backup round-trip keeps member + offset markers ==")
s, xlsx = call("GET", "/api/export/excel", token=TOK, raw=True)
check("export downloads", s == 200 and xlsx[:2] == b"PK", f"{len(xlsx)} bytes")
from openpyxl import load_workbook
wb = load_workbook(io.BytesIO(xlsx))
ws = wb["Transactions"]
header = [c.value for c in ws[1]]
check("export carries the Member and AutoOffset columns",
      "Member" in header and "AutoOffset" in header, str(header))
auto_rows = sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r[header.index("AutoOffset")])
check("the 44 offset rows are exported with their month", auto_rows == 44, str(auto_rows))
before_count = dash["transaction_count"]

import subprocess
import tempfile
path = "round_trip_check.xlsx"
with open(path, "wb") as f:
    f.write(xlsx)
out = subprocess.run(
    ["python", "-c",
     "import json,urllib.request,urllib.error\n"
     "B='http://127.0.0.1:8002'\n"
     f"tok={TOK!r}\n"
     "import csv,io\n"
     "boundary='----x'\n"
     "body=b''\n"
     "def part(name, value):\n"
     "    return (f'--{boundary}\\r\\nContent-Disposition: form-data; name=\"{name}\"\\r\\n\\r\\n{value}\\r\\n').encode()\n"
     f"data=open({path!r},'rb').read()\n"
     "body+=part('confirm','REPLACE_ALL')\n"
     "body+=(f'--{boundary}\\r\\nContent-Disposition: form-data; name=\"file\"; filename=\"rt.xlsx\"\\r\\nContent-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\\r\\n\\r\\n').encode()+data+b'\\r\\n'\n"
     "body+=f'--{boundary}--\\r\\n'.encode()\n"
     "req=urllib.request.Request(B+'/api/admin/import', data=body, method='POST', headers={'Authorization':'Bearer '+tok,'Content-Type':f'multipart/form-data; boundary={boundary}'})\n"
     "try:\n"
     "    r=urllib.request.urlopen(req, timeout=300); print(r.status, r.read().decode())\n"
     "except urllib.error.HTTPError as e: print(e.code, e.read().decode()[:300])\n"],
    capture_output=True, text=True)
print("     import:", out.stdout.strip()[:220])
s, dash2 = call("GET", "/api/dashboard", token=TOK)
check("restored row count matches the export", dash2["transaction_count"] == before_count,
      f"{dash2['transaction_count']} vs {before_count}")
s, summary2 = call("GET", "/api/savings/summary?months=60", token=TOK)
check("offset markers survived the restore",
      sum(1 for m in summary2["months"] if m["closed"] and m["offset_applied"]) == 44,
      str(sum(1 for m in summary2["months"] if m["closed"] and m["offset_applied"])) + " months")
check("savings maths still holds after the restore",
      all(close(m["savings_total"], m["surplus"]) for m in summary2["months"] if m["closed"]))

print("\n== 5. a manual saving adjusts that month's offset ==")
s, cats = call("GET", "/api/categories", token=TOK)
cat_id = cats[0]["id"]
s, entry = call("POST", "/api/savings/entries",
                {"date": "2025-06-10", "amount": 1_000_000, "note": "extra saving"}, token=TOK)
check("saving entry accepted", s == 200 and entry["type"] == "Savings", f"HTTP {s}")
check("saving is attributed to the author", entry.get("member_id") == entry.get("created_by"),
      f"member_id={entry.get('member_id')}")
s, summary3 = call("GET", "/api/savings/summary?months=60", token=TOK)
jun = [m for m in summary3["months"] if m["month"] == "2025-06"][0]
check("the month absorbed the extra saving (offset fell by 1,000,000)",
      close(jun["savings_manual"], 1_000_000) and close(jun["savings_total"], jun["surplus"]),
      f"manual={jun['savings_manual']:,.0f} offset={jun['savings_offset']:,.0f} total={jun['savings_total']:,.0f}")
check("the extra saving is attributed per member",
      close(sum(jun["members"].values()), 1_000_000), str(jun["members"]))

s, _ = call("POST", "/api/savings/entries", {"date": "2025-06-10", "amount": 0}, token=TOK)
check("a zero saving is refused (400)", s == 400, f"HTTP {s}")
s, _ = call("POST", "/api/savings/entries", {"date": "2025-06-10", "amount": 5, "member_id": 99999}, token=TOK)
check("an unknown member is refused (400)", s == 400, f"HTTP {s}")

print("\n== 6. a late expense in a closed month is swept too ==")
s, exp = call("POST", "/api/transactions",
              {"date": "2025-06-15", "type": "Expenses", "category_id": cat_id,
               "amount": 400_000, "description": "late expense"}, token=TOK)
check("late expense added", s == 200, f"HTTP {s}")
s, summary4 = call("GET", "/api/savings/summary?months=60", token=TOK)
jun4 = [m for m in summary4["months"] if m["month"] == "2025-06"][0]
check("2025-06 savings followed the smaller surplus",
      close(jun4["savings_total"], jun4["surplus"]) and close(jun4["savings_total"], jun["savings_total"] - 400_000),
      f"{jun4['savings_total']:,.0f} vs {jun['savings_total'] - 400_000:,.0f}")

print("\n== 7. the running month is left alone ==")
today = time.strftime("%Y-%m-%d")
s, inc = call("POST", "/api/transactions",
              {"date": today, "type": "Income", "category_id": [c for c in cats if c["type"] == "Income"][0]["id"],
               "amount": 9_000_000, "description": "this month income"}, token=TOK)
check("income in the running month added", s == 200, f"HTTP {s}")
s, summary5 = call("GET", "/api/savings/summary?months=60", token=TOK)
cur = [m for m in summary5["months"] if not m["closed"]]
check("running month present but not swept",
      bool(cur) and not cur[0]["offset_applied"] and close(cur[0]["savings_offset"], 0)
      and close(cur[0]["unallocated"], cur[0]["surplus"]),
      f"offset={cur[0]['savings_offset'] if cur else 'n/a'} unallocated={cur[0]['unallocated'] if cur else 'n/a'}")

print("\n== 8. income statement ==")
s, rep = call("GET", "/api/statements/income?start=2022-01-01&end=2026-08-31", token=TOK)
check("income statement loads", s == 200 and "income_lines" in rep, f"HTTP {s}")
check("income and expense lines are present",
      len(rep["income_lines"]) > 0 and len(rep["expense_lines"]) > 0,
      f"{len(rep['income_lines'])} income / {len(rep['expense_lines'])} expense lines")
check("expense groups are subtotalled", len(rep["expense_groups"]) >= 1, f"{len(rep['expense_groups'])} groups")
check("totals are internally consistent",
      close(rep["income_total"] - rep["expense_total"], rep["surplus"]), f"surplus={rep['surplus']:,.0f}")
check("closed months leave nothing unallocated", close(rep["unallocated"], 0, 5),
      f"unallocated={rep['unallocated']:,.0f} (savings={rep['savings_total']:,.0f})")
check("monthly rows are returned", len(rep["months"]) == 44, str(len(rep["months"])))
s, rep_one = call("GET", "/api/statements/income?start=2025-06-01&end=2025-06-30", token=TOK)
check("a single month can be reported", s == 200 and len(rep_one["months"]) == 1,
      f"HTTP {s}, {len(rep_one.get('months', []))} month(s)")

print("\n== 9. balance sheet ==")
s, bs = call("GET", "/api/statements/balance-sheet", token=TOK)
check("balance sheet loads", s == 200 and "net_worth" in bs, f"HTTP {s}")
check("financial assets equal income - expenses",
      close(bs["cash_and_savings"]["total"],
            bs["from_activity"]["income"] - bs["from_activity"]["expenses"], 5),
      f"{bs['cash_and_savings']['total']:,.0f}")
check("net worth equals assets - liabilities",
      close(bs["net_worth"], bs["total_assets"] - bs["liabilities_total"]),
      f"{bs['net_worth']:,.0f}")
base_net = bs["net_worth"]

s, prop = call("POST", "/api/assets",
               {"name": "Family house", "kind": "property", "value": 2_000_000_000,
                "note": "primary residence"}, token=TOK)
check("an asset can be added", s == 201 and prop["kind"] == "property", f"HTTP {s}")
s, car = call("POST", "/api/assets", {"name": "Family car", "kind": "vehicle", "value": 250_000_000}, token=TOK)
check("a vehicle can be added", s == 201, f"HTTP {s}")
s, inv = call("POST", "/api/assets", {"name": "Index funds", "kind": "investment", "value": 100_000_000}, token=TOK)
check("an investment can be added", s == 201, f"HTTP {s}")
s, _ = call("POST", "/api/assets", {"name": "Bad", "kind": "spaceship", "value": 1}, token=TOK)
check("an invalid asset kind is refused (400)", s == 400, f"HTTP {s}")

s, loan = call("POST", "/api/liabilities",
               {"name": "House mortgage", "kind": "mortgage", "outstanding": 750_000_000,
                "monthly_payment": 9_000_000, "interest_rate": 7.5}, token=TOK)
check("a liability can be added", s == 201 and loan["kind"] == "mortgage", f"HTTP {s}")

s, bs2 = call("GET", "/api/statements/balance-sheet", token=TOK)
check("asset items are totalled by kind",
      close(bs2["asset_items_total"], 2_350_000_000),
      f"{bs2['asset_items_total']:,.0f}")
check("liabilities are totalled", close(bs2["liabilities_total"], 750_000_000),
      f"{bs2['liabilities_total']:,.0f}")
check("net worth grew by assets minus liabilities",
      close(bs2["net_worth"], base_net + 1_600_000_000), f"{bs2['net_worth']:,.0f}")
check("asset kinds are offered to the UI",
      "property" in bs2["kinds"]["asset"] and "mortgage" in bs2["kinds"]["liability"],
      str(bs2["kinds"]))

s, upd = call("PUT", f"/api/assets/{car['id']}", {"value": 200_000_000}, token=TOK)
check("an asset value can be updated", s == 200 and close(upd["value"], 200_000_000), f"HTTP {s}")
check("updating kept the other fields", upd["name"] == "Family car" and upd["kind"] == "vehicle", upd["name"])
s, _ = call("DELETE", f"/api/assets/{inv['id']}", token=TOK)
check("an asset can be deleted", s == 200, f"HTTP {s}")
s, _ = call("DELETE", f"/api/liabilities/{loan['id']}", token=TOK)
check("a liability can be deleted", s == 200, f"HTTP {s}")
s, _ = call("DELETE", f"/api/assets/{prop['id']}", token=TOK)
s, _ = call("DELETE", f"/api/assets/{car['id']}", token=TOK)
s, bs3 = call("GET", "/api/statements/balance-sheet", token=TOK)
check("net worth back to its base after cleanup", close(bs3["net_worth"], base_net), f"{bs3['net_worth']:,.0f}")

print("\n== 10. permissions ==")
s, u = call("POST", "/api/users", {"username": f"view{int(time.time())%10000}",
                                   "email": f"view{int(time.time())%10000}@example.com",
                                   "password": "Viewer12345", "role": "viewer",
                                   "is_approved": True, "is_active": True}, token=TOK)
vname = u["username"]
s, vtok = call("POST", "/api/auth/login", {"username": vname, "password": "Viewer12345"})
VTOK = vtok["access_token"]
s, _ = call("POST", "/api/savings/entries", {"date": today, "amount": 1000}, token=VTOK)
check("viewer cannot record a saving (403)", s == 403, f"HTTP {s}")
s, _ = call("POST", "/api/assets", {"name": "x", "kind": "other", "value": 1}, token=VTOK)
check("viewer cannot add an asset (403)", s == 403, f"HTTP {s}")
s, _ = call("POST", "/api/admin/offsets/run", token=VTOK)
check("viewer cannot run the sweep (403)", s == 403, f"HTTP {s}")
s, _ = call("GET", "/api/statements/balance-sheet", token=VTOK)
check("viewer can read the balance sheet", s == 200, f"HTTP {s}")
s, _ = call("GET", "/api/savings/summary", token=VTOK)
check("viewer can read savings", s == 200, f"HTTP {s}")
call("DELETE", f"/api/users/{u['id']}", token=TOK)

print(f"\n===== {passed} passed, {failed} failed =====")
raise SystemExit(1 if failed else 0)
